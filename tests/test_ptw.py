import datetime as dt
import os
import random

import docx
from docx.oxml.ns import qn
import pytest

from ptw import docgen, generator, storage


@pytest.fixture
def store(tmp_path, monkeypatch):
    monkeypatch.setattr(storage, "DATA_DIR", str(tmp_path / "data"))
    st = storage.Store()
    st.history = []
    st.settings["output_dir"] = str(tmp_path / "out")
    return st


def make_plan(store, days=60, seed=1):
    start = dt.date(2026, 1, 1)
    return generator.Planner(store, random.Random(seed)).plan(start, start + dt.timedelta(days=days - 1))


def test_one_permit_per_day(store):
    plan = make_plan(store, 30)
    assert [p["date"] for p in plan] == [(dt.date(2026, 1, 1) + dt.timedelta(d)).isoformat()
                                        for d in range(30)]


def test_job_not_repeated_within_window(store):
    window = store.settings["no_repeat_days"]
    plan = make_plan(store, 90)
    last = {}
    for p in plan:
        d = dt.date.fromisoformat(p["date"])
        if p["job_id"] in last:
            assert (d - last[p["job_id"]]).days > window
        last[p["job_id"]] = d


def test_crew_matches_job_ranks_and_ce_not_used(store):
    for p in make_plan(store, 90):
        job = store.job(p["job_id"])
        ranks = [store.member(i)["rank"] for i in p["crew"]]
        assert store.member(p["lead"])["rank"] in job["lead_ranks"]
        allowed = (set(job["lead_ranks"]) | set(job["officer_ranks"]) | set(job["rating_ranks"])
                   | set(job["pc_ranks"] or store.settings["pc_supervisor_ranks"]))
        assert set(ranks) <= allowed
        assert "CE" not in ranks
        assert len(set(p["crew"])) == len(p["crew"])
        assert p["crew"] == generator.sort_by_rank(store, p["crew"])


def test_only_allowed_ranks_sign_as_pc(store):
    """4E (cadet), 5E and ratings never sign sections 5, 6 or 7 (PSC)."""
    plan = make_plan(store, 120)
    for p in plan:
        assert p["pc"] and p["close"], p
        for role in ("pc", "transfer", "close"):
            if p[role]:
                assert p[role] in p["crew"]
                assert store.member(p[role])["rank"] in ("2E", "3E", "ETO", "AEO"), (role, p)


def test_2e_and_3e_take_turns_as_pc_for_4e_jobs(store):
    plan = make_plan(store, 150)
    sup = [store.member(p["pc"])["rank"] for p in plan
           if store.member(p["lead"])["rank"] in ("4E", "5E") and "remesan" not in p["crew"]]
    assert len(sup) > 30
    assert set(sup) == {"2E", "3E"}
    assert abs(sup.count("2E") - sup.count("3E")) <= len(sup) * 0.25


def test_loto_is_empty_in_generated_permits(store):
    """LOTO is only added by the user (Edit... / Manual), never by the generator."""
    assert all(p["loto"] == "" for p in make_plan(store, 60))


def test_loto_person_when_ticked(store):
    assert generator.pick_loto(store, ["singh", "kaushal"], "singh") == "remesan"  # AEO first
    store.member("remesan")["active"] = False
    assert generator.pick_loto(store, ["singh", "kaushal"], "singh") == "singh"  # then the P.C.
    assert "kaushal" not in generator.loto_candidates(store)  # 4E (cadet) cannot apply LOTO


def test_lead_rotates(store):
    plan = make_plan(store, 120)
    same = sum(a["lead"] == b["lead"] for a, b in zip(plan, plan[1:]))
    assert same < len(plan) * 0.15


def test_crew_spread(store):
    plan = make_plan(store, 120)
    count = {}
    for p in plan:
        for c in p["crew"]:
            count[c] = count.get(c, 0) + 1
    motormen = [count.get(m["id"], 0) for m in store.crew if m["rank"] == "MM"]
    assert min(motormen) > 0.6 * max(motormen)


def test_several_permits_per_day_use_different_people(store):
    store.settings["permits_per_day"] = 2
    plan = make_plan(store, 20)
    assert len(plan) == 40
    for d in {p["date"] for p in plan}:
        a, b = [p for p in plan if p["date"] == d]
        assert a["job_id"] != b["job_id"]


def test_history_influences_next_plan(store):
    first = make_plan(store, 5)
    store.history = first
    nxt = generator.Planner(store, random.Random(2)).plan(dt.date(2026, 1, 6), dt.date(2026, 1, 8))
    assert not {p["job_id"] for p in nxt} & {p["job_id"] for p in first[-3:]}


def test_reroll_keeps_date(store):
    plan = make_plan(store, 7)
    old = plan[3]
    generator.Planner(store, random.Random(5)).reroll(plan, 3)
    assert plan[3]["date"] == old["date"]
    assert plan[3]["job_id"] != old["job_id"]


def _all_text(path):
    """Text of every paragraph once (merged table cells are not repeated)."""
    body = docx.Document(path).element.body
    return "\n".join("".join(t.text or "" for t in p.iter(qn("w:t")))
                     for p in body.iter(qn("w:p")))


def test_write_permits_fills_template(store):
    store.settings["permit_numbering"] = "sequential"
    store.settings["next_permit_no"] = 41
    store.settings["ship_name"] = "MSC Test Ship"
    plan = [{"date": "2026-07-13", "job_id": "me_fuel_valves_overhaul",
             "crew": ["sukhoverkhii", "singh", "kaushal"], "lead": "sukhoverkhii",
             "pc": "singh", "transfer": "sukhoverkhii", "close": "sukhoverkhii",
             "loto": "remesan"}]
    path = docgen.write_permits(store, plan)[0][0]
    assert path.endswith("2026-07-13 ME Fuel Valves Overhaul.docx")
    text = _all_text(path)
    assert "{{" not in text
    for s in ("Ship:  MSC Test Ship", "13/07/2026", "E-   41/2026",
              "Main Engine fuel injection valves",
              "Risk Assessment number RA-001", "2ND ENG DMITRII, 3RD ENG ARUN, 4TH ENG PARTH",
              "Chief Engineer / Vadim Medvedev", "Third Engineer / Arun Singh",
              "Asst. Second Engineer / Dmitrii Sukhoverkhii", "AEO Abhijith Remesan",
              "08:00", "12:00", "20:00"):
        assert s in text, s
    assert store.settings["next_permit_no"] == 42
    assert store.history[-1]["file"] == "2026-07-13 ME Fuel Valves Overhaul.docx"


def test_no_transfer_leaves_section_6_empty(store):
    plan = [{"date": "2026-07-14", "job_id": "fo_strainers_cleaning", "crew": ["singh", "pal"],
             "pc": "singh", "transfer": "", "close": "singh"}]
    text = _all_text(docgen.write_permits(store, plan)[0][0])
    assert text.count("14/07/2026") == 3  # header, section 5, section 7
    assert "12:00" not in text


def test_same_name_gets_suffix(store):
    p = {"date": "2026-07-14", "job_id": "fo_strainers_cleaning", "crew": ["singh"],
         "pc": "singh", "transfer": "", "close": "singh"}
    a, b = docgen.write_permits(store, [dict(p), dict(p)])[0]
    assert a != b and b.endswith("(2).docx")


def test_file_name_is_safe(store):
    store.settings["filename_pattern"] = "{date} {title} / {ra}"
    name = docgen.file_name(store, {"date": "2026-07-14", "job_id": "fo_strainers_cleaning"})
    assert "/" not in name and name.endswith(".docx")


def _runs_with(path, needle):
    body = docx.Document(path).element.body
    for r in body.iter(qn("w:r")):
        text = "".join(t.text or "" for t in r.findall(qn("w:t")))
        if needle in text:
            sz = r.find(qn("w:rPr")).find(qn("w:sz"))
            yield text, int(sz.get(qn("w:val"))) / 2


def test_long_text_is_shrunk_to_one_line(store):
    job = store.job("me_fuel_valves_overhaul")
    job["work"] = ("Removal, overhaul, pressure test and refitting of all fuel valves, "
                   "renewal of nozzles, springs, O-rings and HP pipe gaskets on all units")
    p = {"date": "2026-07-14", "job_id": job["id"], "crew": ["sukhoverkhii"],
         "pc": "sukhoverkhii", "transfer": "", "close": "sukhoverkhii"}
    (path,), warnings = docgen.write_permits(store, [p])
    (text, size), = _runs_with(path, "Removal, overhaul")
    assert text == job["work"] and 6.5 <= size < 9
    assert not warnings
    (_, size), = _runs_with(path, "Main Engine fuel injection valves")
    assert size == 9  # short values keep the normal size


def test_too_long_text_is_cut_and_reported(store):
    job = store.job("me_fuel_valves_overhaul")
    job["work"] = "Very long description " * 20
    p = {"date": "2026-07-14", "job_id": job["id"], "crew": ["sukhoverkhii"],
         "pc": "sukhoverkhii", "transfer": "", "close": "sukhoverkhii"}
    (path,), warnings = docgen.write_permits(store, [p])
    (text, size), = _runs_with(path, "Very long description")
    assert text.endswith("…") and size == 6.5
    assert warnings == [(os.path.basename(path), ["WORK"])]


def test_big_team_uses_compact_ranks(store):
    everybody = [m["id"] for m in store.crew if m["id"] != "ce"]
    p = {"date": "2026-07-14", "job_id": "fo_strainers_cleaning", "crew": everybody,
         "pc": "singh", "transfer": "", "close": "singh"}
    (path,), _ = docgen.write_permits(store, [p])
    text = _all_text(path)
    assert "2E DMITRII, 3E ARUN" in text and "2ND ENG DMITRII" not in text


def test_fit_check_for_editor(store):
    tpl = storage.template_path(store.settings)
    assert docgen.check_fit(tpl, "WORK", "Short text") is None
    size, ok = docgen.check_fit(tpl, "WORK", "x" * 120)
    assert ok and 6.5 <= size < 9
    assert docgen.check_fit(tpl, "WORK", "x" * 400)[1] is False


def test_skip_weekdays(store):
    store.settings["skip_weekdays"] = [5, 6]
    plan = make_plan(store, 14)
    assert len(plan) == 10
    assert all(dt.date.fromisoformat(p["date"]).weekday() < 5 for p in plan)


def test_all_default_jobs_fit_at_normal_size(store):
    tpl = storage.template_path(store.settings)
    for j in store.jobs:
        for key in ("location", "plant", "work", "ra_no"):
            assert docgen.check_fit(tpl, key.upper(), j[key]) is None, (j["id"], key)
    ids = [j["id"] for j in store.jobs]
    assert len(ids) == len(set(ids)) >= 100
    ranks = {r["code"] for r in store.ranks}
    for j in store.jobs:
        assert set(j["lead_ranks"]) | set(j["officer_ranks"]) | set(j["rating_ranks"]) <= ranks


def test_every_ui_string_is_translated():
    tk = pytest.importorskip("tkinter")  # noqa: F841
    import ast
    import re as _re

    from ptw import gui
    from ptw.i18n import RU
    src = open(gui.__file__, encoding="utf-8").read()
    strings = set(gui.WEEKDAYS) | {t for _, t in gui.PLAN_COLUMNS + gui.HistoryTab.COLS}
    for n in ast.walk(ast.parse(src)):
        if isinstance(n, ast.Call) and getattr(n.func, "id", None) == "T" and n.args \
                and isinstance(n.args[0], ast.Constant):
            strings.add(n.args[0].value)
        if isinstance(n, ast.Assign) and getattr(n.targets[0], "id", "") == "hint":
            strings.add(ast.literal_eval(n.value))
    strings |= {m.group(1) for m in _re.finditer(
        r'\("\w+", "([^"]+)", "(?:str|int|bool|codes|range|rank)"\)', src)}
    strings |= {t for _, t in gui.SPLASH_STEPS}
    strings |= {"5) P.C. (accepts permit):", "6) Transferred to:", "7) P.C. (closes permit):"}
    assert [s for s in strings if s not in RU] == []


def test_sound_files_and_player(monkeypatch):
    import wave

    from ptw import sound
    for name in ("startup", "tada", "ding", "chord", "exclamation", "shutdown"):
        with wave.open(sound.path(name)) as w:
            assert w.getnchannels() == 1 and w.getframerate() == 22050
            assert 0.3 < w.getnframes() / w.getframerate() < 3.5
    monkeypatch.setattr(sound, "enabled", True)
    sound.play("does_not_exist")  # missing file: silently nothing
    monkeypatch.setattr(sound, "enabled", False)
    sound.play("startup")


def _numbers(store, dates):
    return generator.permit_numbers(store, [{"date": d} for d in dates])


def test_permit_number_is_day_of_year(store):
    assert store.settings["permit_numbering"] == "day"
    assert _numbers(store, ["2026-10-12", "2026-10-13"]) == ["285", "286"]
    assert _numbers(store, ["2026-10-12", "2026-10-12", "2026-01-05"]) == ["285-1", "285-2", "5"]
    store.history = [{"date": "2026-10-12", "job_id": "x", "crew": [], "permit_no": "285"}]
    assert _numbers(store, ["2026-10-12"]) == ["285-2"]  # second permit that day
    assert generator.permit_numbers(store, [{"date": "2026-10-12"}], replace_dates=True) == ["285"]


def test_day_number_is_printed(store):
    p = {"date": "2026-10-12", "job_id": "fo_strainers_cleaning", "crew": ["singh"],
         "pc": "singh", "transfer": "", "close": "singh"}
    (path,), _ = docgen.write_permits(store, [p])
    assert "E-   285/2026" in _all_text(path)
    assert store.history[-1]["permit_no"] == "285"


def test_blank_numbering(store):
    store.settings["permit_numbering"] = "blank"
    assert _numbers(store, ["2026-10-12"]) == [""]


def _boxes(path):
    rows = []
    for tr in docx.Document(path).element.body.iter(qn("w:tr")):
        boxes = list(tr.iter(qn("w:checkBox")))
        if len(boxes) == 2:
            rows.append(tuple(b.find(qn("w:default")).get(qn("w:val")) for b in boxes))
    return rows


@pytest.mark.parametrize("ao,pc", [(True, True), (True, False), (False, False)])
def test_ao_pc_boxes_ticked_by_setting(store, ao, pc):
    store.settings.update(tick_ao=ao, tick_pc=pc)
    p = {"date": "2026-10-12", "job_id": "fo_strainers_cleaning", "crew": ["singh"],
         "pc": "singh", "transfer": "", "close": "singh"}
    (path,), _ = docgen.write_permits(store, [p])
    rows = _boxes(path)
    assert len(rows) == 12  # 6 hazards + 6 precautions
    assert set(rows) == {("1" if ao else "0", "1" if pc else "0")}


def test_ticks_on_by_default():
    assert storage.DEFAULT_SETTINGS["tick_ao"] and storage.DEFAULT_SETTINGS["tick_pc"]


def test_no_diesel_generator_names_and_responsibilities(store):
    for j in store.jobs:
        assert "DG " not in j["title"] and "Diesel generator" not in j["plant"], j["title"]
        assert "D/G" not in j["location"]
    lead = {j["title"]: j["lead_ranks"][0] for j in store.jobs}
    for t, r in lead.items():
        if t.startswith(("Boiler", "STP", "Sewage")):
            assert r == "3E", t
        if "Purifier" in t or t.startswith(("FO ", "DO ")):
            assert r == "4E", t
    titles = set(lead)
    for t in ("HFO Purifier Bowl Cleaning", "LO Purifier Bowl Cleaning", "DO Purifier Overhaul",
              "ME TC Air Filter Renewal", "AE TC Air Filter Renewal", "Main SW Pump Strainer"):
        assert t in titles


def test_old_settings_and_ranks_are_migrated(tmp_path, monkeypatch):
    import json
    data = tmp_path / "data"
    data.mkdir()
    (data / "settings.json").write_text(json.dumps({"auto_number": True}))
    (data / "ranks.json").write_text(json.dumps([{"code": "4E", "short": "4TH ENG",
                                                   "full": "Fourth Engineer", "officer": True,
                                                   "order": 4},
                                                  {"code": "3E", "short": "3RD ENG",
                                                   "full": "Third Engineer", "officer": True,
                                                   "order": 3}]))
    monkeypatch.setattr(storage, "DATA_DIR", str(data))
    st = storage.Store()
    assert st.settings["permit_numbering"] == "day"
    assert "auto_number" not in st.settings
    assert st.rank("4E")["can_pc"] is False and st.rank("3E")["can_pc"] is True
    assert all("loto" in j for j in st.jobs)


def test_sounds_are_short():
    import wave

    from ptw import sound
    for name in ("startup", "tada", "ding", "chord", "exclamation", "shutdown"):
        with wave.open(sound.path(name)) as w:
            assert w.getnframes() / w.getframerate() <= 1.6, name


def test_icons_exist():
    for size in (16, 32, 48, 64):
        assert os.path.exists(storage.resource("icons", f"ptw_{size}.png"))
    Image = pytest.importorskip("PIL.Image")  # only to read the .ico sizes
    ico = Image.open(storage.resource("icons", "ptw.ico"))
    assert {(16, 16), (32, 32), (256, 256)} <= set(ico.info["sizes"])


def test_upgrade_from_v12_job_list(tmp_path, monkeypatch):
    """A ship running v1.2: DG jobs, old leads, own RA numbers, history with DG jobs."""
    import json
    import shutil
    data = tmp_path / "data"
    data.mkdir()
    old = json.load(open(storage.default_path("upgrade/jobs_v1.2")))
    for j in old:
        if j["id"] == "hfo_purifier_overhaul":
            j["ra_no"] = "ER-PUR-01"           # user's own RA number
        if j["id"] == "dg_fuel_injectors_overhaul":
            j["work"] = "My own text for the DG"  # user's own text
    (data / "jobs.json").write_text(json.dumps(old))
    (data / "history.json").write_text(json.dumps(
        [{"date": "2026-10-01", "job_id": "dg_cylinder_head_overhaul", "crew": ["singh"],
          "pc": "singh", "transfer": "", "close": "singh", "permit_no": "274",
          "file": "x.docx"}]))
    shutil.copy(storage.default_path("crew"), data / "crew.json")
    monkeypatch.setattr(storage, "DATA_DIR", str(data))
    st = storage.Store()
    ids = {j["id"] for j in st.jobs}
    assert not any(i.startswith("dg_") for i in ids)
    assert not any("DG " in j["title"] or "Diesel generator" in j["plant"] for j in st.jobs)
    assert len(st.jobs) == len(storage.default_records("jobs"))
    assert st.job("sewage_plant_maintenance")["lead_ranks"] == ["3E"]
    assert st.job("fo_strainers_cleaning")["lead_ranks"] == ["4E"]
    assert st.job("hfo_purifier_overhaul")["ra_no"] == "ER-PUR-01"     # kept
    assert st.job("ae_fuel_injectors_overhaul")["work"] == "My own text for the AE"
    assert st.job("hfo_purifier_bowl_cleaning")["group"] == "hfo_purifier"
    assert st.history[0]["job_id"] == "ae_cylinder_head_overhaul"
    assert (data / "jobs_backup_v1.2.json").exists()
    assert st.settings["data_version"] == storage.DATA_VERSION
    # second start: nothing changes any more
    st.jobs = [j for j in st.jobs if j["id"] != "ae_piston_overhaul"]
    storage.save("jobs", st.jobs)
    assert storage.Store().job("ae_piston_overhaul") is None


def test_old_template_and_sounds_are_replaced(tmp_path, monkeypatch):
    """Files shipped by an older version (known md5) are replaced, the user's own are kept.

    No git needed: "old" files are made here and their md5 is declared as shipped."""
    import shutil
    app = tmp_path / "app"
    (app / "templates").mkdir(parents=True)
    (app / "sounds").mkdir()
    repo = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    old_tpl = app / "templates" / "ptw_template.docx"
    d = docx.Document(os.path.join(repo, "templates", "ptw_template.docx"))
    d.add_paragraph("old version of the template")
    d.save(str(old_tpl))
    old_snd = app / "sounds" / "startup.wav"
    old_snd.write_bytes(b"RIFF old startup sound of v1.2")
    own_snd = app / "sounds" / "ding.wav"
    own_snd.write_bytes(b"RIFF my own sound")
    monkeypatch.setattr(storage, "OLD_TEMPLATE_MD5", {storage.md5(str(old_tpl))})
    monkeypatch.setattr(storage, "OLD_SOUND_MD5", {storage.md5(str(old_snd))})
    monkeypatch.setattr(storage, "APP_DIR", str(app))
    monkeypatch.setattr(storage, "BUNDLE_DIR", repo)
    monkeypatch.setattr(storage, "DATA_DIR", str(app / "data"))
    monkeypatch.setattr(storage, "DEFAULT_TEMPLATE", str(old_tpl))
    storage.ensure_data()
    assert storage.md5(str(old_tpl)) == storage.md5(os.path.join(repo, "templates", "ptw_template.docx"))
    assert (app / "templates" / "ptw_template_old.docx").exists()      # backup of the old one
    assert storage.md5(str(old_snd)) == storage.md5(os.path.join(repo, "sounds", "startup.wav"))
    assert own_snd.read_bytes() == b"RIFF my own sound"                # user's own kept
    # a template the user changed himself is never replaced
    shutil.copy(app / "templates" / "ptw_template_old.docx", old_tpl)
    with open(old_tpl, "ab") as f:
        f.write(b" ")
    before = storage.md5(str(old_tpl))
    storage.ensure_data()
    assert storage.md5(str(old_tpl)) == before


def test_printed_dates_do_not_update(store):
    """Word must not replace the permit date by the computer date when printing."""
    p = {"date": "2026-03-05", "job_id": "fo_strainers_cleaning", "crew": ["singh"],
         "pc": "singh", "transfer": "singh", "close": "singh"}
    (path,), _ = docgen.write_permits(store, [p])
    body = docx.Document(path).element.body
    instr = " ".join(x.text or "" for x in body.iter(qn("w:instrText")))
    assert "DATE" not in instr and "TIME" not in instr
    types = {t.get(qn("w:val")) for t in body.iter(qn("w:type"))}
    assert not types & {"currentDate", "currentTime"}
    assert _all_text(path).count("05/03/2026") == 4
    # an old template with the date fields is cleaned while filling, too
    old = docx.Document(storage.template_path(store.settings))
    assert docgen.freeze_auto_fields(old) == 0  # shipped template is already clean


def test_numbers_never_repeat_after_history_delete(store):
    p = {"date": "2026-10-12", "job_id": "fo_strainers_cleaning", "crew": ["singh"],
         "pc": "singh", "transfer": "", "close": "singh"}
    docgen.write_permits(store, [dict(p), dict(p)])
    assert [h["permit_no"] for h in store.history] == ["285-1", "285-2"]
    del store.history[0]
    docgen.write_permits(store, [dict(p)])
    assert [h["permit_no"] for h in store.history] == ["285-2", "285-3"]


def test_permit_without_pc_is_refused(store):
    p = {"date": "2026-10-12", "job_id": "fo_strainers_cleaning", "crew": ["kaushal", "padilla"],
         "pc": "", "transfer": "", "close": ""}
    with pytest.raises(docgen.PermitError):
        docgen.write_permits(store, [p])
    assert store.history == []


def test_pc_falls_back_when_supervisor_ranks_not_on_board(store):
    store.settings["pc_supervisor_ranks"] = ["XX"]
    for p in make_plan(store, 40):
        assert p["pc"] and store.can_pc(p["pc"])


def test_4e_job_with_aeo_helper_gets_2e_or_3e_as_pc(store):
    job = store.job("ac_plant_maintenance")
    pl = generator.Planner(store, random.Random(1))
    for k in range(30):
        t = pl.make_team(job, dt.date(2026, 5, 1) + dt.timedelta(days=k))
        if store.member(t["lead"])["rank"] in ("4E", "5E"):
            assert store.member(t["pc"])["rank"] in ("2E", "3E")
            assert store.member(t["close"])["rank"] in ("2E", "3E")


def test_grouped_jobs_not_close_together(store):
    plan = make_plan(store, 200)
    last = {}
    for p in plan:
        g = store.job(p["job_id"])["group"]
        if g:
            d = dt.date.fromisoformat(p["date"])
            if g in last:
                assert (d - last[g]).days > store.settings["no_repeat_days"], (g, p["date"])
            last[g] = d


def test_reroll_respects_repeat_window_both_ways(store):
    plan = make_plan(store, 30)
    window = store.settings["no_repeat_days"]
    pl = generator.Planner(store, random.Random(9))
    for i in range(0, 30, 3):
        pl.reroll(plan, i)
        d = dt.date.fromisoformat(plan[i]["date"])
        for q in plan:
            if q is not plan[i] and q["job_id"] == plan[i]["job_id"]:
                assert abs((dt.date.fromisoformat(q["date"]) - d).days) > window


def test_replace_overwrites_files_instead_of_adding_copies(store):
    p = {"date": "2026-10-12", "job_id": "fo_strainers_cleaning", "crew": ["singh"],
         "pc": "singh", "transfer": "", "close": "singh"}
    q = {"date": "2026-10-13", "job_id": "boiler_burner_maintenance", "crew": ["singh"],
         "pc": "singh", "transfer": "", "close": "singh"}
    docgen.write_permits(store, [dict(p), dict(q)])
    folder = store.settings["output_dir"]
    first = sorted(os.listdir(folder))
    docgen.write_permits(store, [dict(p), dict(q)], replace_dates=True)
    assert sorted(os.listdir(folder)) == first  # same names, no "(2)"
    assert len(store.history) == 2
    docgen.write_permits(store, [dict(p)])      # "keep both": a second file
    assert len(os.listdir(folder)) == 3
    # another job on a replaced day: the old file of that day goes
    r = dict(p, job_id="sludge_pump_strainer")
    docgen.write_permits(store, [r], replace_dates=True)
    names = os.listdir(folder)
    assert not any("FO Strainers" in n for n in names) and any("Sludge" in n for n in names)
