"""Filling the Word template and saving permits."""
import datetime as dt
import io
import os
import re

import docx
from docx.oxml.ns import qn

from . import fontmetrics, generator, storage

PLACEHOLDER = re.compile(r"\{\{\s*([A-Z_]+)\s*\}\}")
_XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


def _paragraphs(document):
    parts = [document.element.body]
    for s in document.sections:
        for hf in (s.header, s.footer, s.first_page_header, s.first_page_footer,
                   s.even_page_header, s.even_page_footer):
            if not hf.is_linked_to_previous:
                parts.append(hf._element)
    for part in parts:
        yield from list(part.iter(qn("w:p")))


def _replace_in_paragraph(p, values):
    """Replace {{KEY}} even when Word split it over several runs.

    Returns the keys found and the w:t elements that received a value."""
    ts = list(p.iter(qn("w:t")))
    full = "".join(t.text or "" for t in ts)
    if "{{" not in full:
        return [], []
    # character offset where each w:t starts
    starts, pos = [], 0
    for t in ts:
        starts.append(pos)
        pos += len(t.text or "")
    keys, filled = [], []
    for m in reversed(list(PLACEHOLDER.finditer(full))):
        key = m.group(1)
        if key not in values:
            continue
        keys.append(key)
        a, b = m.span()
        first = True
        for t, s in zip(ts, starts):
            text = t.text or ""
            e = s + len(text)
            if e <= a or s >= b:
                continue
            lo, hi = max(a, s) - s, min(b, e) - s
            t.text = text[:lo] + (str(values[key]) if first else "") + text[hi:]
            t.set(_XML_SPACE, "preserve")
            if first:
                filled.append(t)
            first = False
    return keys, filled


# ---------- keeping every value on one line (the form must fit one page) ----------

DEFAULT_CELL_MARGIN = 108  # dxa (twentieths of a point), Word default left/right
# leave a little room: Word and LibreOffice lay out text slightly differently
SAFETY_RATIO, SAFETY_PT = 0.99, 3.0


# font sizes from the document styles (a run without w:sz is NOT 9 pt: Word uses the
# paragraph / character style, then docDefaults, then 10 pt)
_SIZES = {"default": 10.0, "para": {}, "char": {}, "para_default": None}


def _prepare_sizes(document):
    st = document.styles.element
    default = 10.0
    dd = st.find(qn("w:docDefaults"))
    if dd is not None:
        sz = dd.find(f"{qn('w:rPrDefault')}/{qn('w:rPr')}/{qn('w:sz')}")
        if sz is not None:
            default = int(sz.get(qn("w:val"))) / 2
    raw = {}
    for el in st.findall(qn("w:style")):
        sid = el.get(qn("w:styleId"))
        based = el.find(qn("w:basedOn"))
        rpr = el.find(qn("w:rPr"))
        sz = rpr.find(qn("w:sz")) if rpr is not None else None
        raw[sid] = (el.get(qn("w:type")), based.get(qn("w:val")) if based is not None else None,
                    int(sz.get(qn("w:val"))) / 2 if sz is not None else None,
                    el.get(qn("w:default")) == "1")

    def resolve(sid, seen=()):
        if sid not in raw or sid in seen:
            return None
        _t, based, size, _d = raw[sid]
        return size if size is not None else resolve(based, seen + (sid,))

    _SIZES["default"] = default
    _SIZES["para"] = {sid: resolve(sid) for sid, v in raw.items() if v[0] == "paragraph"}
    _SIZES["char"] = {sid: resolve(sid) for sid, v in raw.items() if v[0] == "character"}
    _SIZES["para_default"] = next((sid for sid, v in raw.items()
                                   if v[0] == "paragraph" and v[3]), None)


def _run_size(r):
    rpr = r.find(qn("w:rPr"))
    if rpr is not None:
        sz = rpr.find(qn("w:sz"))
        if sz is not None:
            return int(sz.get(qn("w:val"))) / 2
        rs = rpr.find(qn("w:rStyle"))
        if rs is not None and _SIZES["char"].get(rs.get(qn("w:val"))):
            return _SIZES["char"][rs.get(qn("w:val"))]
    p = r.getparent()
    while p is not None and p.tag != qn("w:p"):
        p = p.getparent()
    sid = _SIZES["para_default"]
    if p is not None:
        ps = p.find(f"{qn('w:pPr')}/{qn('w:pStyle')}")
        if ps is not None:
            sid = ps.get(qn("w:val"))
    return _SIZES["para"].get(sid) or _SIZES["default"]


def _run_bold(r):
    rpr = r.find(qn("w:rPr"))
    b = rpr.find(qn("w:b")) if rpr is not None else None
    return b is not None and b.get(qn("w:val"), "true") not in ("0", "false")


# children of w:rPr that must come after w:sz / w:szCs (ECMA-376 CT_RPr order)
_AFTER_SZ = ("szCs", "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign", "rtl",
             "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath")


def _set_run_size(r, size):
    rpr = r.find(qn("w:rPr"))
    if rpr is None:
        rpr = r.makeelement(qn("w:rPr"), {})
        r.insert(0, rpr)
    half = str(int(round(size * 2)))
    for tag in ("sz", "szCs"):
        el = rpr.find(qn("w:" + tag))
        if el is None:
            el = rpr.makeelement(qn("w:" + tag), {})
            later = _AFTER_SZ[_AFTER_SZ.index(tag) + 1:] if tag == "szCs" else _AFTER_SZ
            nxt = next((c for c in rpr if c.tag in {qn("w:" + t) for t in later}), None)
            if nxt is not None:
                nxt.addprevious(el)
            else:
                rpr.append(el)
        el.set(qn("w:val"), half)


def _grid_width(tc):
    """Width of the table grid columns the cell spans (dxa), or None."""
    tr = tc.getparent()
    tbl = tr.getparent() if tr is not None else None
    grid = tbl.find(qn("w:tblGrid")) if tbl is not None else None
    if grid is None:
        return None
    cols = [int(g.get(qn("w:w"), 0)) for g in grid.findall(qn("w:gridCol"))]
    trpr = tr.find(qn("w:trPr"))
    gb = trpr.find(qn("w:gridBefore")) if trpr is not None else None
    idx = int(gb.get(qn("w:val"))) if gb is not None else 0
    for c in tr.findall(qn("w:tc")):
        span = c.find(f"{qn('w:tcPr')}/{qn('w:gridSpan')}")
        n = int(span.get(qn("w:val"))) if span is not None else 1
        if c is tc:
            return sum(cols[idx:idx + n]) or None
        idx += n
    return None


def _available_width(p):
    """Width in points a paragraph inside a table cell can use, or None."""
    tc = p.getparent()
    while tc is not None and tc.tag != qn("w:tc"):
        tc = tc.getparent()
    if tc is None:
        return None
    tcpr = tc.find(qn("w:tcPr"))
    w = tcpr.find(qn("w:tcW")) if tcpr is not None else None
    if w is None or w.get(qn("w:type")) not in ("dxa", None):
        return None
    width = int(w.get(qn("w:w")))
    grid = _grid_width(tc)
    if grid:  # auto-layout tables are laid out on the grid, which can be narrower
        width = min(width, grid)
    width -= 2 * DEFAULT_CELL_MARGIN
    ppr = p.find(qn("w:pPr"))
    ind = ppr.find(qn("w:ind")) if ppr is not None else None
    if ind is not None:
        for side in ("left", "start", "right", "end", "firstLine"):
            v = ind.get(qn("w:" + side))
            if v:
                width -= int(v)
    return width / 20 * SAFETY_RATIO - SAFETY_PT


def _value_runs(p, filled):
    """Runs holding inserted values, in document order."""
    wanted = {t.getparent() for t in filled}
    return [r for r in p.iter(qn("w:r")) if r in wanted]


def _cut(runs, room, size):
    """Shorten the text of runs so that it is at most `room` points wide."""
    ell = fontmetrics.text_width("\u2026", size)
    used, cutting = 0.0, False
    for r in runs:
        for t in r.findall(qn("w:t")):
            text = t.text or ""
            if cutting:
                t.text = ""
                continue
            keep = ""
            for ch in text:
                w = fontmetrics.text_width(ch, size, _run_bold(r))
                if used + w > room - ell:
                    cutting = True
                    break
                keep += ch
                used += w
            if cutting:
                t.text = keep.rstrip(" ,") + "\u2026"
            else:
                t.text = keep


def fit_paragraph(p, filled, min_size):
    """Keep the paragraph on one line: shrink the inserted values, cut them as a last resort.

    Returns "ok", "shrunk" or "cut"."""
    avail = _available_width(p)
    if avail is None or not filled:
        return "ok"
    runs = list(p.iter(qn("w:r")))
    if any(r.find(qn("w:tab")) is not None for r in runs):
        return "ok"  # tab stops: cannot measure reliably
    vruns = _value_runs(p, filled)
    fixed = value = 0.0
    for r in runs:
        text = "".join(t.text or "" for t in r.findall(qn("w:t")))
        w = fontmetrics.text_width(text, _run_size(r), _run_bold(r))
        if r in vruns:
            value += w
        else:
            fixed += w
    if fixed + value <= avail or value == 0:
        return "ok"
    scale = (avail - fixed) / value
    status = "shrunk"
    sizes = []
    for r in vruns:
        size = int(_run_size(r) * scale * 2) / 2  # Word uses half points
        if size < min_size:
            size, status = min_size, "cut"
        sizes.append(size)
        _set_run_size(r, size)
    if status == "cut":
        _cut(vruns, avail - fixed, min(sizes))
    return status


AUTO_FIELDS = ("DATE", "TIME", "PRINTDATE", "SAVEDATE", "CREATEDATE", "EDITTIME")
AUTO_INPUTS = ("currentDate", "currentTime")


def _fields(p):
    """Complex fields of a paragraph: dicts with the runs of begin / separate / end,
    the field instruction and the depth (0 = outermost)."""
    out, stack = [], []
    for r in p.iter(qn("w:r")):
        fc = r.find(qn("w:fldChar"))
        it = r.find(qn("w:instrText"))
        if fc is not None:
            kind = fc.get(qn("w:fldCharType"))
            if kind == "begin":
                stack.append({"begin": r, "sep": None, "instr": "", "depth": len(stack),
                              "ffdata": fc.find(qn("w:ffData")), "runs": [r]})
                continue
            if stack:
                f = stack[-1]
                f["runs"].append(r)
                if kind == "separate":
                    f["sep"] = r
                elif kind == "end":
                    f["end"] = r
                    out.append(stack.pop())
                    if stack:
                        stack[-1]["runs"].extend(out[-1]["runs"])
            continue
        if stack:
            stack[-1]["runs"].append(r)
            if it is not None and stack[-1]["sep"] is None:
                stack[-1]["instr"] += it.text or ""
    return out


def _is_auto(field):
    """Fields Word recalculates before printing: DATE / TIME ... and date form fields."""
    words = field["instr"].split()
    if words and words[0].upper() in AUTO_FIELDS:
        return True
    ff = field["ffdata"]
    if ff is not None:
        t = ff.find(qn("w:textInput") + "/" + qn("w:type"))
        if t is not None and t.get(qn("w:val")) in AUTO_INPUTS:
            return True
    return False


def freeze_auto_fields(document):
    """Turn self-updating date/time fields into plain text.

    The MSC form has its dates in 'current date' form fields: Word showed our date
    on screen but printed the date of the computer. Keeping only the field result
    makes the printed date the permit date. Returns the number of fields frozen."""
    frozen = 0
    for p in _paragraphs(document):
        for _ in range(10):
            fields = [f for f in _fields(p) if _is_auto(f) and "end" in f]
            if not fields:
                break
            f = min(fields, key=lambda f: f["depth"])  # outermost first
            runs, sep = f["runs"], f["sep"]
            keep = False
            for r in runs:
                if r is sep:
                    keep = True
                    r.getparent().remove(r)
                    continue
                if r is f["end"]:
                    r.getparent().remove(r)
                    break
                if not keep or sep is None:
                    if r.getparent() is not None:
                        r.getparent().remove(r)
            frozen += 1
    return frozen


def set_checkboxes(document, tick_ao, tick_pc):
    """Tick the A.O. / P.C. form check boxes of sections 2 and 3.

    Every table row with exactly two check boxes is a hazard / precaution line:
    the first box is the A.O. column, the second the P.C. column."""
    for tr in document.element.body.iter(qn("w:tr")):
        boxes = list(tr.iter(qn("w:checkBox")))
        if len(boxes) != 2:
            continue
        for box, on in zip(boxes, (tick_ao, tick_pc)):
            val = "1" if on else "0"
            default = box.find(qn("w:default"))
            if default is None:  # order: size / sizeAuto, default, checked
                default = box.makeelement(qn("w:default"), {})
                size = box.find(qn("w:size"))
                if size is None:
                    size = box.find(qn("w:sizeAuto"))
                if size is not None:
                    size.addnext(default)
                else:
                    box.insert(0, default)
            default.set(qn("w:val"), val)
            checked = box.find(qn("w:checked"))
            if checked is None:
                checked = box.makeelement(qn("w:checked"), {})
                default.addnext(checked)
            checked.set(qn("w:val"), val)


_template_cache = {}


def _template_bytes(template):
    """The template is read once (a network drive can be slow), again only if it changed."""
    stamp = (template, os.path.getmtime(template), os.path.getsize(template))
    data = _template_cache.get(template)
    if not data or data[0] != stamp:
        with open(template, "rb") as f:
            data = (stamp, f.read())
        _template_cache[template] = data
    return data[1]


def fill_template(template, values, out_path, fit=True, min_size=6.5, ticks=(True, True)):
    """Fill the template; returns the keys whose text had to be cut to fit one line."""
    d = docx.Document(io.BytesIO(_template_bytes(template)))
    _prepare_sizes(d)
    too_long = []
    for p in _paragraphs(d):
        keys, filled = _replace_in_paragraph(p, values)
        if fit and filled and fit_paragraph(p, filled, min_size) == "cut":
            too_long += keys
    if ticks is not None:
        set_checkboxes(d, *ticks)
    freeze_auto_fields(d)
    d.save(out_path)
    return too_long


_limits_cache = {}


def field_limits(template):
    """{KEY: (free width in pt, font size)} for placeholders that sit in table cells."""
    try:
        stamp = (template, os.path.getmtime(template))
    except OSError:
        return {}
    if stamp in _limits_cache:
        return _limits_cache[stamp]
    limits = {}
    d = docx.Document(template)
    _prepare_sizes(d)
    for p in _paragraphs(d):
        avail = _available_width(p)
        if avail is None:
            continue
        fixed, size, keys = 0.0, _SIZES["default"], []
        for r in p.iter(qn("w:r")):
            text = "".join(t.text or "" for t in r.findall(qn("w:t")))
            found = PLACEHOLDER.findall(text)
            if found:
                keys += found
                size = _run_size(r)
            fixed += fontmetrics.text_width(PLACEHOLDER.sub("", text), _run_size(r), _run_bold(r))
        if len(keys) == 1:  # a field used twice (A.O.): the narrower place counts
            old = limits.get(keys[0])
            if old is None or avail - fixed < old[0]:
                limits[keys[0]] = (avail - fixed, size)
    _limits_cache[stamp] = limits
    return limits


def check_fit(template, key, text, min_size=6.5):
    """None if text fits at normal size, else (size it will be printed at, fits at all)."""
    lim = field_limits(template).get(key)
    if not lim or not text:
        return None
    free, size = lim
    width = fontmetrics.text_width(text, size)
    if width <= free:
        return None
    new = int(size * free / width * 2) / 2
    return max(new, min_size), new >= min_size


REQUIRED_PLACEHOLDERS = ("SHIP", "DATE", "PERMIT_NO", "LOCATION", "PLANT", "WORK",
                         "ISSUED_TO", "RA_NO", "LOTO_NAME", "AO", "PC", "PC_CLOSE")


def template_missing(template):
    """Placeholders the template lacks (an old v1.2 template has no SHIP / LOTO_NAME)."""
    try:
        d = docx.Document(template)
    except Exception:  # noqa: BLE001 - reported to the user as a missing template
        return list(REQUIRED_PLACEHOLDERS)
    text = "".join(t.text or "" for p in _paragraphs(d) for t in p.iter(qn("w:t")))
    found = set(PLACEHOLDER.findall(text))
    return [k for k in REQUIRED_PLACEHOLDERS if k not in found]


def values_for(store, permit):
    s = store.settings
    day = generator.parse_date(permit["date"])
    job = store.job(permit["job_id"]) or {}
    no = permit.get("permit_no")
    empty_field = "\u2002" * 5  # how Word shows an empty form field
    return {
        "SHIP": s.get("ship_name", ""),
        "LOTO_NAME": generator.loto_text(store, permit.get("loto")) or empty_field,
        "DATE": day.strftime("%d/%m/%Y"),
        "YEAR": day.strftime("%Y"),
        "PERMIT_NO": str(no) if no else " " * 14,
        "TIME_ISSUE": s["time_issue"],
        "TIME_TRANSFER": s["time_transfer"] if permit.get("transfer") else "",
        "TRANSFER_DATE": day.strftime("%d/%m/%Y") if permit.get("transfer") else "",
        "TIME_CLOSE": s["time_close"],
        "LOCATION": job.get("location", ""),
        "PLANT": job.get("plant", ""),
        "WORK": job.get("work", ""),
        "RA_NO": job.get("ra_no", ""),
        "ISSUED_TO": generator.issued_to_text(store, permit["crew"]),
        "AO": generator.sign_text(store, s.get("ao_id")),
        "PC": generator.sign_text(store, permit.get("pc")),
        "TRANSFER_TO": generator.sign_text(store, permit.get("transfer")),
        "PC_CLOSE": generator.sign_text(store, permit.get("close")),
    }


_BAD = re.compile(r'[<>:"/\\|?*\x00-\x1f]')


def file_name(store, permit):
    s = store.settings
    job = store.job(permit["job_id"]) or {"title": permit["job_id"]}
    day = generator.parse_date(permit["date"])
    name = s.get("filename_pattern") or "{date} {title}"
    name = name.format(date=day.strftime(s.get("filename_date_format") or "%Y-%m-%d"),
                       title=job["title"], ra=job.get("ra_no", ""),
                       no=permit.get("permit_no") or "")
    name = _BAD.sub("_", name).strip(" .") or "permit"
    return name + ".docx"


def unique_path(folder, name):
    path = os.path.join(folder, name)
    base, ext = os.path.splitext(path)
    n = 2
    while os.path.exists(path):
        path = f"{base} ({n}){ext}"
        n += 1
    return path


class PermitError(ValueError):
    """Permits that must not be written (no P.C., unknown job ...)."""


def _check_all(store, permits):
    bad = []
    for p in permits:
        problems = generator.check_permit(store, p)
        if problems:
            job = store.job(p.get("job_id")) or {"title": p.get("job_id")}
            bad.append(f"{p.get('date')}  {job['title']}: {', '.join(problems)}")
    if bad:
        raise PermitError("\n".join(bad))


def _prepare(store):
    s = store.settings
    folder = s.get("output_dir")
    if not folder:
        raise ValueError("Output folder is not set (Settings).")
    template = storage.template_path(s)
    if not os.path.exists(template):
        raise FileNotFoundError(f"Template not found: {template}")
    os.makedirs(folder, exist_ok=True)
    return folder, template


def _write_one(store, p, path, template):
    """Fill one permit into path; returns the keys cut to keep one line."""
    s = store.settings
    min_size = float(s.get("min_font_size", 6.5))
    values = values_for(store, p)
    res = check_fit(template, "ISSUED_TO", values["ISSUED_TO"], min_size)
    if res and not res[1]:  # a big team: "2E DMITRII" instead of "2ND ENG DMITRII"
        values["ISSUED_TO"] = generator.issued_to_text(store, p["crew"], compact=True)
    return fill_template(template, values, path, fit=s.get("fit_one_line", True),
                         min_size=min_size,
                         ticks=(bool(s.get("tick_ao", True)), bool(s.get("tick_pc", True))))


def write_permits(store, permits, replace_dates=False, progress=None, cancelled=None):
    """Create .docx files, number them and record history.

    progress(done, total, file name) is called after every file; cancelled() -> True
    stops after the current file (the files written so far stay in the history).
    Returns (paths, warnings); warnings name permits whose text had to be cut.
    Raises PermitError (nothing written) if a permit has no P.C., an unknown job ..."""
    s = store.settings
    _check_all(store, permits)
    folder, template = _prepare(store)

    numbers = generator.permit_numbers(store, permits, replace_dates)
    if replace_dates:
        # "replace": the old permits of these days go - history entries AND their files,
        # so the new files get the plain names instead of "... (2).docx"
        dates = {p["date"] for p in permits}
        old = [h for h in store.history if h["date"] in dates]
        old_files = [os.path.join(folder, h["file"]) for h in old if h.get("file")]
        old_files = [f for f in old_files if os.path.isfile(f)]
        for f in old_files:  # check first: a file open in Word cannot be replaced
            try:
                with open(f, "r+b"):
                    pass
            except OSError as e:
                raise PermitError(f"{os.path.basename(f)}: {e.strerror or e} "
                                  "(close it in Word and try again)") from None
        for f in old_files:
            os.remove(f)
        store.history = [h for h in store.history if h["date"] not in dates]

    paths, warnings = [], []
    try:
        for k, (p, no) in enumerate(zip(permits, numbers), 1):
            if cancelled and cancelled():
                break
            if not p.get("permit_no") and no:
                p["permit_no"] = no
                if s.get("permit_numbering") == "sequential" and str(no).isdigit():
                    s["next_permit_no"] = int(no) + 1
            path = unique_path(folder, file_name(store, p))
            too_long = _write_one(store, p, path, template)
            paths.append(path)
            if too_long:
                warnings.append((os.path.basename(path), sorted(set(too_long))))
            store.history.append({**p, "file": os.path.basename(path),
                                  "created": dt.datetime.now().isoformat(timespec="seconds")})
            if progress:
                progress(k, len(permits), os.path.basename(path))
    finally:  # also after an error / cancel: the files already written are recorded
        store.history.sort(key=lambda h: h["date"])
        storage.save("history", store.history)
        storage.save("settings", store.settings)
    return paths, warnings


def rewrite_files(store, entries, progress=None, cancelled=None):
    """Write the files of history entries again (same number and file name, overwritten),
    e.g. permits made with an old template. Returns (paths, warnings)."""
    folder, template = _prepare(store)
    paths, warnings = [], []
    for k, h in enumerate(entries, 1):
        if cancelled and cancelled():
            break
        p = dict(h)
        if not store.job(p.get("job_id")):
            raise PermitError(f"{p.get('date')}: job {p.get('job_id')} is not in the list")
        path = os.path.join(folder, h.get("file") or file_name(store, p))
        too_long = _write_one(store, p, path, template)
        paths.append(path)
        if too_long:
            warnings.append((os.path.basename(path), sorted(set(too_long))))
        if progress:
            progress(k, len(entries), os.path.basename(path))
    return paths, warnings
